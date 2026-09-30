import io
import re
import uuid
from dataclasses import dataclass, field

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.user import User
from app.repositories import audit_repository, group_repository, role_repository, user_repository
from app.schemas.user import ImportConfirmResponse, ImportPreviewResponse, ImportRowResult

TEMPLATE_HEADERS = ["full_name", "email", "role", "groups"]
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class ParsedRow:
    row_number: int
    full_name: str | None
    email: str | None
    role_name: str | None
    group_names: list[str]
    errors: list[str] = field(default_factory=list)
    role_id: uuid.UUID | None = None
    group_ids: list[uuid.UUID] = field(default_factory=list)

    def to_result(self) -> ImportRowResult:
        return ImportRowResult(
            row_number=self.row_number,
            full_name=self.full_name,
            email=self.email,
            role=self.role_name,
            groups=self.group_names,
            errors=self.errors,
        )


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "users"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["Jane Doe", "jane.doe@example.com", "USER", "New hires"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _parse_workbook(file_bytes: bytes) -> list[ParsedRow]:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises varied exceptions for bad files
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The file is not a valid Excel file.") from exc

    ws = wb.active
    rows: list[ParsedRow] = []
    for idx, raw_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if raw_row is None or all(cell in (None, "") for cell in raw_row):
            continue
        full_name = str(raw_row[0]).strip() if len(raw_row) > 0 and raw_row[0] else None
        email = str(raw_row[1]).strip().lower() if len(raw_row) > 1 and raw_row[1] else None
        role_name = str(raw_row[2]).strip().upper() if len(raw_row) > 2 and raw_row[2] else None
        groups_raw = str(raw_row[3]).strip() if len(raw_row) > 3 and raw_row[3] else ""
        group_names = [g.strip() for g in groups_raw.split(",") if g.strip()]
        rows.append(
            ParsedRow(row_number=idx, full_name=full_name, email=email, role_name=role_name, group_names=group_names)
        )
    return rows


def validate_rows(db: Session, file_bytes: bytes) -> list[ParsedRow]:
    rows = _parse_workbook(file_bytes)

    seen_emails: dict[str, int] = {}
    for row in rows:
        if not row.full_name:
            row.errors.append("Full name is required.")
        if not row.email:
            row.errors.append("Email is required.")
        elif not EMAIL_PATTERN.match(row.email):
            row.errors.append("The email format is not valid.")
        elif row.email in seen_emails:
            row.errors.append(f"Duplicated email in the file (row {seen_emails[row.email]}).")
        else:
            seen_emails[row.email] = row.row_number

    valid_emails = [row.email for row in rows if row.email and EMAIL_PATTERN.match(row.email)]
    existing = user_repository.existing_emails(db, valid_emails)
    for row in rows:
        if row.email in existing:
            row.errors.append("A user with that email already exists on the platform.")

    for row in rows:
        if not row.role_name:
            row.errors.append("Role is required.")
        else:
            role = role_repository.get_by_name(db, row.role_name)
            if role is None:
                row.errors.append(f"Invalid role: {row.role_name}.")
            else:
                row.role_id = role.id

        for group_name in row.group_names:
            group = group_repository.get_by_name(db, group_name)
            if group is None:
                row.errors.append(f"Group not found: {group_name}.")
            else:
                row.group_ids.append(group.id)

    return rows


def preview_import(db: Session, file_bytes: bytes) -> ImportPreviewResponse:
    rows = validate_rows(db, file_bytes)
    valid = sum(1 for r in rows if not r.errors)
    return ImportPreviewResponse(
        total_rows=len(rows),
        valid_rows=valid,
        invalid_rows=len(rows) - valid,
        rows=[r.to_result() for r in rows],
    )


def confirm_import(db: Session, *, current_user: User, file_bytes: bytes) -> ImportConfirmResponse:
    rows = validate_rows(db, file_bytes)

    created = 0
    for row in rows:
        if row.errors:
            continue
        user = user_repository.create(db, full_name=row.full_name, email=row.email, role_id=row.role_id)
        db.flush()
        if row.group_ids:
            user_repository.set_groups(db, user_id=user.id, group_ids=row.group_ids)
        audit_repository.log_action(
            db,
            user_id=current_user.id,
            action="USER_CREATED",
            entity_type="user",
            entity_id=user.id,
            new_data={"full_name": row.full_name, "email": row.email, "source": "excel_import"},
        )
        created += 1

    db.commit()
    return ImportConfirmResponse(
        created=created,
        skipped=len(rows) - created,
        rows=[r.to_result() for r in rows],
    )

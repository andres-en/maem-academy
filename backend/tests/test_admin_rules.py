"""Administration rules: users, bulk import, groups, categories and uploaded files."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from tests.helpers import create_course


def _role_id(client, headers, name: str) -> str:
    return next(r["id"] for r in client.get("/api/v1/roles", headers=headers).json() if r["name"] == name)


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["full_name", "email", "role", "groups"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Users --------------------------------------------------------------------


def test_admin_creates_a_user_and_emails_are_unique(client, make_user):
    _, admin = make_user("ADMIN")
    payload = {"full_name": "Jane Doe", "email": "jane@example.com", "role_id": _role_id(client, admin, "USER")}
    created = client.post("/api/v1/users", json=payload, headers=admin)
    assert created.status_code == 201
    assert created.json()["status"] == "ACTIVE"

    duplicate = client.post("/api/v1/users", json=payload, headers=admin)
    assert duplicate.status_code == 409
    assert duplicate.json()["detail"] == "A user with that email already exists."


def test_users_must_be_deactivated_before_deletion_and_cannot_delete_themselves(client, make_user):
    admin_id, admin = make_user("ADMIN")
    user_id, _ = make_user("USER")

    blocked = client.delete(f"/api/v1/users/{user_id}", headers=admin)
    assert blocked.status_code == 400
    assert blocked.json()["detail"] == "You must deactivate the user before deleting it."

    assert (
        client.patch(f"/api/v1/users/{user_id}/status", json={"status": "INACTIVE"}, headers=admin).status_code == 200
    )
    assert client.delete(f"/api/v1/users/{user_id}", headers=admin).status_code == 204

    myself = client.delete(f"/api/v1/users/{admin_id}", headers=admin)
    assert myself.status_code == 400
    assert myself.json()["detail"] == "You cannot delete your own user."


def test_force_delete_is_superadmin_only_and_requires_typing_the_email(client, make_user):
    _, admin = make_user("ADMIN")
    _, superadmin = make_user("SUPERADMIN")
    user_id, _ = make_user("USER", email="gone@example.com")
    client.patch(f"/api/v1/users/{user_id}/status", json={"status": "INACTIVE"}, headers=superadmin)

    assert (
        client.request(
            "DELETE", f"/api/v1/users/{user_id}/force", json={"confirm_email": "gone@example.com"}, headers=admin
        ).status_code
        == 403
    )
    wrong = client.request(
        "DELETE", f"/api/v1/users/{user_id}/force", json={"confirm_email": "nope@example.com"}, headers=superadmin
    )
    assert wrong.status_code == 400
    ok = client.request(
        "DELETE", f"/api/v1/users/{user_id}/force", json={"confirm_email": "gone@example.com"}, headers=superadmin
    )
    assert ok.status_code == 204


# --- Bulk import ------------------------------------------------------------------


def test_import_template_has_the_expected_columns(client):
    resp = client.get("/api/v1/users/import/template")
    assert resp.status_code == 200
    header = next(load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True))
    assert list(header) == ["full_name", "email", "role", "groups"]


def test_bulk_import_validates_every_row_before_creating_anything(client, make_user):
    _, admin = make_user("ADMIN")
    make_user("USER", email="taken@example.com")
    client.post("/api/v1/groups", json={"name": "Sales"}, headers=admin)
    file = _xlsx(
        [
            ["Ana Lopez", "ana@example.com", "USER", "Sales"],
            ["No Email", None, "USER", ""],
            ["Bad Email", "not-an-email", "USER", ""],
            ["Taken", "taken@example.com", "USER", ""],
            ["Bad Role", "role@example.com", "WIZARD", ""],
            ["Dup", "ana@example.com", "USER", ""],
            ["Unknown Group", "grp@example.com", "USER", "Ghosts"],
        ]
    )
    files = {"file": ("users.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = client.post("/api/v1/users/import/validate", files=files, headers=admin).json()
    assert (preview["total_rows"], preview["valid_rows"], preview["invalid_rows"]) == (7, 1, 6)
    errors = {row["row_number"]: row["errors"] for row in preview["rows"]}
    assert errors[2] == []
    assert errors[3] == ["Email is required."]
    assert errors[4] == ["The email format is not valid."]
    assert errors[5] == ["A user with that email already exists on the platform."]
    assert errors[6] == ["Invalid role: WIZARD."]
    assert any("Duplicated email in the file" in e for e in errors[7])
    assert errors[8] == ["Group not found: Ghosts."]

    confirmed = client.post("/api/v1/users/import/confirm", files=files, headers=admin).json()
    assert (confirmed["created"], confirmed["skipped"]) == (1, 6)
    users = client.get("/api/v1/users", params={"search": "ana@example.com"}, headers=admin).json()
    assert users["total"] == 1


def test_bulk_import_rejects_files_that_are_not_excel(client, make_user):
    _, admin = make_user("ADMIN")
    files = {"file": ("users.xlsx", b"definitely not a spreadsheet", "application/octet-stream")}
    resp = client.post("/api/v1/users/import/validate", files=files, headers=admin)
    assert resp.status_code == 400
    assert resp.json()["detail"] == "The file is not a valid Excel file."


# --- Groups and categories --------------------------------------------------------


def test_group_membership_add_and_remove(client, make_user):
    _, admin = make_user("ADMIN")
    user_id, _ = make_user("USER")
    group = client.post("/api/v1/groups", json={"name": "Support"}, headers=admin).json()

    client.post(f"/api/v1/groups/{group['id']}/users", json={"user_ids": [str(user_id)]}, headers=admin)
    members = client.get(f"/api/v1/groups/{group['id']}", headers=admin).json()["members"]
    assert [m["id"] for m in members] == [str(user_id)]

    assert client.delete(f"/api/v1/groups/{group['id']}/users/{user_id}", headers=admin).status_code == 200
    missing = client.delete(f"/api/v1/groups/{group['id']}/users/{user_id}", headers=admin)
    assert missing.status_code == 404


def test_category_rules(client, make_user):
    _, admin = make_user("ADMIN")
    general = next(c for c in client.get("/api/v1/categories", headers=admin).json() if c["name"] == "GENERAL")
    protected = client.delete(f"/api/v1/categories/{general['id']}", headers=admin)
    assert protected.status_code == 400
    assert protected.json()["detail"] == "The GENERAL category cannot be deleted."

    category = client.post("/api/v1/categories", json={"name": "Compliance"}, headers=admin).json()
    assert client.post("/api/v1/categories", json={"name": "Compliance"}, headers=admin).status_code == 409

    active = client.delete(f"/api/v1/categories/{category['id']}", headers=admin)
    assert active.json()["detail"] == "You must deactivate the category before deleting it."
    client.patch(f"/api/v1/categories/{category['id']}/status", json={"status": "INACTIVE"}, headers=admin)
    assert client.delete(f"/api/v1/categories/{category['id']}", headers=admin).status_code == 204


def test_a_category_used_by_a_course_cannot_be_deleted(client, make_user):
    _, admin = make_user("ADMIN")
    category = client.post("/api/v1/categories", json={"name": "Onboarding"}, headers=admin).json()
    create_course(client, admin, category_ids=[category["id"]])
    client.patch(f"/api/v1/categories/{category['id']}/status", json={"status": "INACTIVE"}, headers=admin)
    resp = client.delete(f"/api/v1/categories/{category['id']}", headers=admin)
    assert resp.status_code == 409


# --- Uploaded files ---------------------------------------------------------------


def test_uploads_are_checked_by_content_not_only_by_extension(client, make_user, storage):
    _, instructor = make_user("INSTRUCTOR")
    course = create_course(client, instructor)
    module = client.post(f"/api/v1/courses/{course['id']}/modules", json={"title": "M"}, headers=instructor).json()
    lesson = client.post(f"/api/v1/course-modules/{module['id']}/lessons", json={"title": "L"}, headers=instructor)
    url = f"/api/v1/lessons/{lesson.json()['id']}/contents"

    fake_pdf = {"file": ("slides.pdf", b"MZ\x90\x00 an executable in disguise", "application/pdf")}
    rejected = client.post(url, data={"content_type": "PDF"}, files=fake_pdf, headers=instructor)
    assert rejected.status_code == 400
    assert rejected.json()["detail"] == "The file content does not match the declared type."

    wrong_ext = {"file": ("slides.exe", b"%PDF-1.7 ...", "application/pdf")}
    assert client.post(url, data={"content_type": "PDF"}, files=wrong_ext, headers=instructor).status_code == 400

    real_pdf = {"file": ("slides.pdf", b"%PDF-1.7 minimal", "application/pdf")}
    accepted = client.post(url, data={"content_type": "PDF"}, files=real_pdf, headers=instructor)
    assert accepted.status_code == 201, accepted.text
    assert accepted.json()["file_url"].startswith("http://storage.test/")
    assert list(storage.objects.values()) == [b"%PDF-1.7 minimal"]
